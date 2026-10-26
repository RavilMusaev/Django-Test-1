import inspect
import textwrap
import unittest
from datetime import datetime

from django.core.management.base import BaseCommand
from django.test.runner import DiscoverRunner
from docx import Document
from docx.shared import Pt


def flatten(suite):
    """Разворачивает вложенные TestSuite в плоский список тестов."""
    for item in suite:
        if isinstance(item, unittest.TestSuite):
            yield from flatten(item)
        else:
            yield item


class Command(BaseCommand):
    help = 'Запускает тесты проекта и записывает их код и результаты в Word-файл'

    def add_arguments(self, parser):
        parser.add_argument('--output', default='tests_report.docx')
        parser.add_argument('labels', nargs='*', default=['blog', 'crud'])

    def handle(self, *args, **options):
        runner = DiscoverRunner(verbosity=0, interactive=False)
        runner.setup_test_environment()
        suite = runner.build_suite(options['labels'])
        tests = list(flatten(suite))

        db_config = runner.setup_databases()
        try:
            result = unittest.TestResult()
            suite.run(result)
        finally:
            runner.teardown_databases(db_config)
            runner.teardown_test_environment()

        bad_ids = {t.id() for t, _ in result.failures + result.errors}

        doc = Document()
        doc.add_heading('Отчёт по тестам проекта django-test', level=0)
        doc.add_paragraph(f'Дата: {datetime.now():%d.%m.%Y %H:%M}')
        doc.add_paragraph(
            f'Всего тестов: {len(tests)}. '
            f'Успешно: {len(tests) - len(bad_ids)}. '
            f'С ошибками: {len(bad_ids)}.'
        )

        current_class = None
        for test in sorted(tests, key=lambda t: t.id()):
            cls = type(test)
            if cls is not current_class:
                current_class = cls
                doc.add_heading(f'{cls.__module__}.{cls.__name__}', level=1)

            method = getattr(cls, test._testMethodName)
            doc.add_heading(test._testMethodName, level=2)
            doc.add_paragraph(
                'Результат: ' + ('FAIL' if test.id() in bad_ids else 'OK')
            )

            source = textwrap.dedent(inspect.getsource(method))
            run = doc.add_paragraph().add_run(source)
            run.font.name = 'Consolas'
            run.font.size = Pt(9)

        doc.save(options['output'])
        self.stdout.write(self.style.SUCCESS(
            f'Готово: {options["output"]} ({len(tests)} тестов)'
        ))